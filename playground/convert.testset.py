# -*- coding: utf-8 -*-
import os
import shutil
import subprocess
import threading
import traceback
from queue import Queue
from concurrent.futures import ThreadPoolExecutor, as_completed

import diskcache
from openpyxl import load_workbook
from tqdm import tqdm

# ===================== TOP CONFIG =====================

# 并行 task 数
PARALLEL_DOWNLOAD_TASKS = 1
PARALLEL_CVT_TASKS = 1

# 每个 task 内部资源
DOWNLOAD_WORKERS_PER_TASK = 4
CVT_WORKERS_PER_TASK = 4

# 背压：最多下载领先多少个 task
MAX_AHEAD = 4
SENTINEL = object()

# 路径相关
HOME = "/root/jyw_ws/lerobot_jpeg/playground"

RCLONE_BIN = f"{HOME}/rclone-v1.68.2-linux-amd64/rclone"
XLSX_IN = f"{HOME}/数据需求总表同步_转换列表.xlsx"

LOCAL_TMP_ROOT = "/mnt/fastdisk/tmp_download"
OUTPUT_ROOT = "/mnt/sdc/output_v3_test_0331"
CACHE_DIR = "/root/_v3_test_0331_ckpt"

# 远端相关
REMOTE_ROOT = "huawei-cloud:openloong-apps-prod-private/data-collector-svc/align"

# Excel 列名
TASK_ID_COLUMN = "任务ID"
TASK_DSCP_COLUMN = "任务步骤描述-人话英文版"
DEVICE_COLUMN = "采集设备"

# 设备 -> robot name (used by base_converter)
DEVICE_TO_ROBOT = {
    "星尘智能S1": "astribot_s1",
    "松灵Aloha": "cobotmagic",
    "傅利叶GR-2": "gr2",
    "Franka FR3": "fr3",
}

# hard code test set
HARDCODED_TEST_TASKS = {
    "星尘智能S1": [
        "248fa0186c534f409cc998085477a935",
        "466bf34b5fcc4da3a0977897edea5ad4",
        "18f28bff645e4b62a1ff501928455798",
        "6e100419b1cb40dea8f0879b32acebda",
        "9dd789ac46314d6784a26027a1c3eaf8",
    ],
    "松灵Aloha": [
        "aef6825ceb35485587a3efed4d9401d4",
        "41e5d6faddaa44c7b5be7b18d6d63858",
        "5893af6503dd4ebc9ebf319d35190526",
        "c10826940c3444a099f2995d03fb3f18",
        "9518202cf08641979e349000171452b0",
    ],
    "傅利叶GR-2": [
        "167b536f67db4de998f9ee4118611e54",
        "478f93055aaa48029d5e0f3b3be9c061",
        "6e275d421e394e03a3dd5cc0783867ef",
        "092926d5e59840a2935efc3a2515f293",
        "fd53bb497c464ccaa2a703bbd324ca77",
    ],
    "Franka FR3": [
        "afdad85d6e6a47a8bfa72177d435940f",
        "5d3a1efd0c354fb4bace42847d7f1096",
        "80113d2163114e2d854130c14e69aa5a",
        "0f25153abb0d4101a25f155b8ce9b25d",
        "63bd0530d34d465fb75b9a783bb4533a",
    ],
}

# Base converter script path
BASE_CONVERTER = f"{HOME}/convert/base_converter.py"

# ======================================================
ALLOWED_DEVICES = set(DEVICE_TO_ROBOT.keys())


def run(cmd, check=True):
    print("[RUN]", " ".join(map(str, cmd)))
    p = subprocess.Popen(cmd, start_new_session=False)
    ret = p.wait()
    if check and ret != 0:
        raise subprocess.CalledProcessError(ret, cmd)
    return ret


def get_robot_name(device_name: str) -> str:
    name = str(device_name).strip()
    if name not in DEVICE_TO_ROBOT:
        raise ValueError(f"Unsupported device: {device_name}")
    return DEVICE_TO_ROBOT[name]


def download_task(task_id: str, local_task_dir: str, workers: int):
    remote_task_root = f"{REMOTE_ROOT}/{task_id}"
    os.makedirs(local_task_dir, exist_ok=True)

    cmd = [
        RCLONE_BIN, "copy",
        remote_task_root,
        local_task_dir,
    ]
    run(cmd, check=True)


def cvt_task(task_id: str, local_task_dir: str, task_dscp: str, device_name: str, workers: int, cache: diskcache.Cache):
    robot_name = get_robot_name(device_name)
    output_dir = f"{OUTPUT_ROOT}/{task_id}"
    os.makedirs(os.path.dirname(output_dir), exist_ok=True)

    cmd = [
        "python", BASE_CONVERTER,
        "--robot", robot_name,
        "--input", local_task_dir,
        "--output", output_dir,
        "--repo_id", task_id,
        "--task", str(task_dscp) if task_dscp is not None else "",
        "--workers", str(workers),
    ]
    print(f"device_name: {device_name}, robot_name: {robot_name}")
    print(f"task_dscp: {task_dscp}")

    try:
        run(cmd, check=True)
    except subprocess.CalledProcessError as e:
        cache.set(f"convert_failure_{task_id}", {
            "task_id": task_id,
            "device": device_name,
            "robot": robot_name,
            "cmd": cmd,
            "returncode": e.returncode,
            "traceback": traceback.format_exc(),
        })
        print(f"[FAILURE CACHED] {task_id}")


def cleanup(local_task_dir: str):
    shutil.rmtree(local_task_dir, ignore_errors=True)


def load_selected_tasks():
    wb = load_workbook(XLSX_IN, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    col_idx = {str(name).strip(): i for i, name in enumerate(headers) if name is not None}

    required_cols = [TASK_ID_COLUMN, TASK_DSCP_COLUMN, DEVICE_COLUMN]
    for c in required_cols:
        if c not in col_idx:
            raise KeyError(f"Excel missing required column: {c}")

    task_id_col = col_idx[TASK_ID_COLUMN]
    task_dscp_col = col_idx[TASK_DSCP_COLUMN]
    device_col = col_idx[DEVICE_COLUMN]

    # task_id -> info
    task_info = {}
    for row in ws.iter_rows(min_row=2):
        task_id_val = row[task_id_col].value
        if task_id_val is None:
            continue

        task_id = str(task_id_val).strip()
        if not task_id:
            continue

        device_name = row[device_col].value
        device_name = str(device_name).strip() if device_name is not None else ""

        task_dscp = row[task_dscp_col].value
        task_info[task_id] = {
            "task_id": task_id,
            "device_name": device_name,
            "task_dscp": task_dscp,
        }

    wb.close()

    task_rows = []
    print("=" * 120)
    print("HARDCODED TEST TASKS")
    print("=" * 120)

    for device_name, task_ids in HARDCODED_TEST_TASKS.items():
        robot_name = get_robot_name(device_name)
        print(f"\nDEVICE: {device_name}  ROBOT: {robot_name}")

        for idx, task_id in enumerate(task_ids, 1):
            if task_id not in task_info:
                print(f"  [MISSING_IN_XLSX] {task_id}")
                continue

            info = task_info[task_id]

            if info["device_name"] != device_name:
                print(
                    f"  [DEVICE_MISMATCH] task_id={task_id} "
                    f"xlsx_device={info['device_name']} hardcoded_device={device_name}"
                )
                continue

            task_rows.append((task_id, info["task_dscp"], device_name))
            print(f"{idx:>2}. {task_id}")

    print("\n" + "=" * 120)
    print(f"TOTAL SELECTED TASKS: {len(task_rows)}")
    print("=" * 120)

    return task_rows


def main():
    cache = diskcache.Cache(CACHE_DIR)
    os.makedirs(LOCAL_TMP_ROOT, exist_ok=True)
    os.makedirs(OUTPUT_ROOT, exist_ok=True)

    task_rows = load_selected_tasks()

    ready_q = Queue()
    slots = Queue(maxsize=PARALLEL_CVT_TASKS + MAX_AHEAD)

    def cvt_and_cleanup(task_id: str, local_dir: str, task_dscp: str, device_name: str, cvt_threads: int):
        try:
            cvt_task(task_id, local_dir, task_dscp, device_name, workers=cvt_threads, cache=cache)
            cache.set("cached_task_id", task_id)
            cleanup(local_dir)
            print(f"[DONE] CVT {task_id}")
        except Exception:
            print(traceback.format_exc())
            os._exit(1)
        finally:
            slots.get()
            slots.task_done()
        return task_id

    def cvt_submitter():
        while True:
            item = ready_q.get()
            try:
                if item is SENTINEL:
                    return
                task_id, local_dir, task_dscp, device_name = item
                print("cvt get", task_id, device_name)
                cvt_and_cleanup(task_id, local_dir, task_dscp, device_name, CVT_WORKERS_PER_TASK)
            finally:
                ready_q.task_done()

    def submit_download(dl_pool: ThreadPoolExecutor, task_id: str, task_dscp: str, device_name: str):
        def _job():
            local_dir = os.path.join(LOCAL_TMP_ROOT, task_id)
            print(f"\n=== DOWNLOAD TASK {task_id} ({device_name}) ===")
            download_task(task_id, local_dir, workers=DOWNLOAD_WORKERS_PER_TASK)
            print(f"\n=== DOWNLOAD TASK FINISH {task_id} ({device_name}) ===")
            ready_q.put((task_id, local_dir, task_dscp, device_name))
        return dl_pool.submit(_job)

    cached_task_id = cache.get("cached_task_id")
    skip_mode = cached_task_id is not None
    seen_cached = False

    with ThreadPoolExecutor(max_workers=PARALLEL_DOWNLOAD_TASKS) as dl_pool:
        t = threading.Thread(target=cvt_submitter, daemon=True)
        t.start()

        download_futures = []

        for task_id, task_dscp, device_name in tqdm(task_rows, total=len(task_rows), miniters=1):
            if skip_mode:
                if not seen_cached:
                    if task_id == cached_task_id:
                        seen_cached = True
                    continue
                else:
                    if task_id == cached_task_id:
                        continue
                    skip_mode = False
                    cached_task_id = None

            print("ready_q size:", ready_q.qsize())
            print("slots size:", slots.qsize())

            slots.put(1)
            download_futures.append(submit_download(dl_pool, task_id, task_dscp, device_name))

        for fut in as_completed(download_futures):
            fut.result()

        ready_q.join()
        ready_q.put(SENTINEL)
        t.join()

    cache.close()


if __name__ == "__main__":
    main()