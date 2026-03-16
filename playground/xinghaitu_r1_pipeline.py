# -*- coding: utf-8 -*-
import re
from collections import defaultdict
from openpyxl import load_workbook

XLSX_IN = "task_and_description_checked.xlsx"
DOWNLOAD_LIST_TXT = "download_list.txt"
OUTPUT_TXT = "xinghaitu_r1_filtered_download_list.txt"

# 你要的 robot type（按需改成你的值；可多个）
ROBOT_TYPES = {"星海图R1"}  # 例子：把这里换成你表里 D 列的实际内容

# ---- 1) 从 download_list 建立：data_id -> lines 映射 ----
# 兼容 data_id 是 32位 hex：.../<data_id>/<episode>/...
HEX32 = re.compile(r"^[0-9a-f]{32}$", re.IGNORECASE)

id_to_lines = defaultdict(list)

with open(DOWNLOAD_LIST_TXT, "r", encoding="utf-8") as f:
    for line in f:
        s = line.strip()
        if not s:
            continue
        parts = s.rstrip("/").split("/")
        # 期望最后一段是 episode，倒数第二段是 data_id
        if len(parts) >= 2 :#and HEX32.match(parts[-2]):
            data_id = parts[-2].lower()
            # print(data_id)
            id_to_lines[data_id].append(s)

print("Loaded download_list ids:", len(id_to_lines))


# ---- 2) 从 Excel 过滤 robot type，得到 data_id 列表 ----
wb = load_workbook(XLSX_IN)
ws = wb.active

selected_ids = []
for r in range(2, ws.max_row + 1):
    data_id = ws[f"A{r}"].value

    # if data_id in id_to_lines.keys():
    #     ...
    # else:
    #     print(f"{data_id} not found")

    # continue


    robot_class = ws[f"D{r}"].value  # 你说的 robot type

    if not data_id or not robot_class:
        continue

    robot_class_s = str(robot_class).strip()
    if robot_class_s not in ROBOT_TYPES:
        continue

    data_id_s = str(data_id).strip().lower()
    selected_ids.append(data_id_s)
    # break




selected_ids = list(dict.fromkeys(selected_ids))  # 去重且保序
print("Selected task ids:", len(selected_ids))

# ---- 3) 输出对应行 ----
out_lines = []
missing = []

for data_id in selected_ids:
    lines = id_to_lines.get(data_id)
    if not lines:
        missing.append(data_id)
        continue
    out_lines.extend(lines)

with open(OUTPUT_TXT, "w", encoding="utf-8") as f:
    for s in out_lines:
        f.write(s + "\n")

print("Saved:", OUTPUT_TXT)
print("Total lines:", len(out_lines))
print("Missing ids:", len(missing))
if missing:
    print("Example missing:", missing[:10])