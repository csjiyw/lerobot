# -*- coding: utf-8 -*-
import os
import shutil
import subprocess
from openpyxl import load_workbook
from tqdm import tqdm

HOME="/root/jyw_ws/lerobot/playground"
# ====== CONFIG ======
DOWNLOAD_TASK_TXT = f"{HOME}/filtered_download_list.txt"
RCLONE_BIN = f"{HOME}/rclone-v1.68.2-linux-amd64/rclone"
LOCAL_TMP_ROOT = "/mnt/sdc/tmp_download"

CVT_CMD_TEMPLATE = [
    "python3", "/root/jyw_ws/lerobot/playground/ginie1_align2lerobot_v30.py",
    "--input", "{local_task_dir}",
    "--output", "/mnt/sdc/output/{task_id}",
    "--repo_id", "{task_id}",
    "--task","{task_dscp}"
]

MAX_EPISODES_PER_TASK = 999999
# =====================

XLSX_IN  = f"{HOME}/task_and_description_checked.xlsx"

wb = load_workbook(XLSX_IN)
ws = wb.active  # 或 wb["Sheet1"]

def run(cmd, check=True):
    print("CMD:", " ".join(cmd) if isinstance(cmd, list) else cmd)
    return subprocess.run(cmd, check=check)

def parse_task_and_episode(remote_line: str):
    """
    支持：
      .../<task_id>/<episode_id>/      (目录)
      .../<task_id>/<episode_id>.h5   (文件)
    返回 (task_id, episode_name)
    """
    p = remote_line.strip()
    if not p:
        return None, None
    p2 = p.rstrip("/")
    parts = p2.split("/")
    if len(parts) < 2:
        return None, None
    task_id = parts[-2].strip().lower()
    episode_name = parts[-1].strip()  # 目录名 或 *.h5
    return task_id, episode_name

def download_one(remote_ep: str, task_id: str, episode_name: str):
    local_task_dir = os.path.join(LOCAL_TMP_ROOT, task_id)
    os.makedirs(local_task_dir, exist_ok=True)


    is_file = episode_name.lower().endswith(".h5")

    if is_file:
        local_dest = local_task_dir
    else:
        local_dest = os.path.join(local_task_dir, episode_name)

    cmd = [
        RCLONE_BIN, "copy",
        remote_ep,
        local_dest,
        "--transfers", "8",
        "--checkers", "16",
        "--progress",
    ]
    run(cmd, check=True)
    return local_task_dir

def get_task_root(remote_ep: str) -> str:
    """
    remote_ep:
      .../<task_id>/<episode_id>/      或 .../<task_id>/<episode_id>.h5
    return:
      .../<task_id>
    """
    p = remote_ep.strip().rstrip("/")
    parts = p.split("/")
    if len(parts) < 2:
        raise ValueError(f"bad remote_ep: {remote_ep}")
    return "/".join(parts[:-1])  # 去掉 episode 那层

def download_batch(remote_task_root: str, task_id: str, episode_names: list):
    """
    一次 rclone copy remote_task_root -> local_task_dir
    只下载 episode_names 指定的条目（文件或目录）
    """
    local_task_dir = os.path.join(LOCAL_TMP_ROOT, task_id)
    os.makedirs(local_task_dir, exist_ok=True)

    cmd = [
        RCLONE_BIN, "copy",
        remote_task_root,
        local_task_dir,
        "--transfers", "8",
        "--checkers", "16",
        "--progress", # 并发时建议别用 --progress（输出会炸）
        "--stats", "1s",
    ]

    # include 文件 / 目录
    for ep in episode_names:
        ep = ep.strip().rstrip("/")
        if not ep:
            continue
        if ep.lower().endswith(".h5"):
            # 文件
            cmd += ["--filter", f"+ {ep}"]
        else:
            # 目录
            cmd += ["--filter", f"+ {ep}/**"]

    # 排除其它所有
    cmd += ["--filter", "- **"]

    run(cmd, check=True)
    return local_task_dir



def cvt_task(task_id: str, local_task_dir: str):
    for row in ws.iter_rows(min_row=2):  # 跳过表头
        a_val = str(row[0].value).strip() if row[0].value else ""
        if a_val.lower() == task_id.lower():
            task_dscp = row[10].value   # K 列 = index 10
            break
    print("task_dscp",task_dscp)
    cmd = [x.format(task_id=task_id, local_task_dir=local_task_dir,task_dscp=task_dscp) for x in CVT_CMD_TEMPLATE]
    run(cmd, check=True)

def cleanup(local_task_dir: str):
    shutil.rmtree(local_task_dir, ignore_errors=True)

def main():
    os.makedirs(LOCAL_TMP_ROOT, exist_ok=True)

    cur_task = None
    cur_dir = None
    cur_cnt = 0

    cur_task_root = None
    cur_eps = []   # episode_name list

    

    def flush():
        nonlocal cur_task, cur_dir, cur_cnt, cur_task_root, cur_eps
        if not cur_task or not cur_eps:
            cur_task, cur_dir, cur_cnt, cur_task_root, cur_eps = None, None, 0, None, []
            return
        


        print(f"\n=== DOWNLOAD TASK {cur_task} episodes={len(cur_eps)} ===")
        download_batch(cur_task_root, cur_task, cur_eps)

        # print(f"\n=== CVT TASK {cur_task} episodes={cur_cnt} ===")
      
        # cvt_task(cur_task, cur_dir)
  
        # cleanup(cur_dir)
        # print(f"[DONE] cleaned {cur_dir}")

        cur_task, cur_dir, cur_cnt = None, None, 0
    
    with open(DOWNLOAD_TASK_TXT, "r", encoding="utf-8") as f:
        total_lines = sum(1 for _ in f)

    with open(DOWNLOAD_TASK_TXT, "r", encoding="utf-8") as f:
        for line in tqdm(f, total=total_lines):
            s = line.strip()
            if not s:
                continue

            task_id, ep_name = parse_task_and_episode(s)
            if not task_id:
                continue

            # task 切换：先处理上一组
            if cur_task is not None and task_id != cur_task:
                flush()
                # break

            # 初始化当前 task（并清残留）
            if cur_task is None:
                cur_task = task_id
                cur_dir = os.path.join(LOCAL_TMP_ROOT, cur_task)
                cur_cnt = 0
                cur_eps = []
                cur_task_root = get_task_root(s)

                # if os.path.exists(cur_dir):
                #     shutil.rmtree(cur_dir, ignore_errors=True)
                # os.makedirs(cur_dir, exist_ok=True)

            # if cur_cnt >= MAX_EPISODES_PER_TASK:
            #     continue

            # 下载这一条

            # print(s,cur_task,ep_name)
      
            # download_one(s, cur_task, ep_name)
            cur_eps.append(ep_name)

            cur_cnt += 1
            # break
    

    flush()

if __name__ == "__main__":
    main()