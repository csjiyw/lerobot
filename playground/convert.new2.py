# -*- coding: utf-8 -*-
import os
import shutil
import subprocess
import sys
import threading
import traceback
from queue import Queue
from concurrent.futures import ThreadPoolExecutor, as_completed

import diskcache
from openpyxl import load_workbook
from tqdm import tqdm

# ===================== TOP CONFIG =====================

# 并行 task 数
PARALLEL_DOWNLOAD_TASKS = 1
PARALLEL_CVT_TASKS = 1

# 每个 task 内部资源
DOWNLOAD_WORKERS_PER_TASK = 4
CVT_WORKERS_PER_TASK = 6

# 背压：最多下载领先多少个 task
MAX_AHEAD = 1
SENTINEL = object()

# 路径相关
HOME = "/root/jyw_ws/lerobot_jpeg/playground"

RCLONE_BIN = f"{HOME}/rclone-v1.68.2-linux-amd64/rclone"
XLSX_IN = f"{HOME}/数据需求总表同步_转换列表.xlsx"

LOCAL_TMP_ROOT = "/mnt/fastdisk/tmp_download"
OUTPUT_ROOT = "/mnt/sdc/output_0331"
CACHE_DIR = "/root/_robot3_ckpt"

# 远端相关
REMOTE_ROOT = "huawei-cloud:openloong-apps-prod-private/data-collector-svc/align"

# Excel 列名
TASK_ID_COLUMN = "任务ID"
TASK_DSCP_COLUMN = "任务步骤描述-人话英文版"
DEVICE_COLUMN = "采集设备"

# 设备 -> robot name (used by base_converter)
DEVICE_TO_ROBOT = {
    "松灵Aloha": "cobotmagic",
    "Franka FR3": "fr3",
    "傅利叶GR-2": "gr2",
    "星尘智能S1": "astribot_s1",
}

# Tasks to filter for
TARGET_TASKS = {
    '7fdc9a1d008c4aa186a997485d198cec', 
    '97f2a59c2fef489cb6c7ec7591961cd0',
    '6a09fa0e3e7f470283528c3d44ae899d', 
    '58c29ff274d84adea51424776fb6b92c', 
    'c414d7dfb49d4c5dbe41a179ae645284',
}

# Base converter script path
BASE_CONVERTER = f"{HOME}/convert/base_converter.py"


# ======================================================
ALLOWED_DEVICES = set(DEVICE_TO_ROBOT.keys())

def run(cmd, check=True):
    p = subprocess.Popen(cmd, start_new_session=False)
    ret = p.wait()
    if check and ret != 0:
        raise subprocess.CalledProcessError(ret, cmd)


def get_robot_name(device_name: str) -> str:
    name = str(device_name).strip()
    if name not in DEVICE_TO_ROBOT:
        raise ValueError(f"Unsupported device: {device_name}")
    return DEVICE_TO_ROBOT[name]


def download_task(task_id: str, local_task_dir: str, workers: int):
    remote_task_root = f"{REMOTE_ROOT}/{task_id}"
    os.makedirs(local_task_dir, exist_ok=True)

    cmd = [
        RCLONE_BIN, "copy",
        remote_task_root,
        local_task_dir,
    ]
    run(cmd, check=True)


def cvt_task(task_id: str, local_task_dir: str, task_dscp: str, device_name: str, workers: int, cache: diskcache.Cache):
    output_dir = f"{OUTPUT_ROOT}/{task_id}"
    robot_name = get_robot_name(device_name)

    cmd = [
        "python", BASE_CONVERTER,
        "--robot", robot_name,
        "--input", local_task_dir,
        "--output", output_dir,
        "--repo_id", task_id,
        "--task", str(task_dscp) if task_dscp is not None else "",
        "--workers", str(workers),
    ]
    print(f"device_name: {device_name}, robot_name: {robot_name}")
    print(f"task_dscp: {task_dscp}")

    # Inline subprocess call to capture stderr without crashing the pipeline
    result = subprocess.run(cmd, stderr=subprocess.PIPE, text=True)
    
    if result.returncode != 0:
        stderr_msg = result.stderr if result.stderr else "No stderr captured"
        cache.set(f"convert_failure_{task_id}", {
            "task_id": task_id,
            "device": device_name,
            "robot": robot_name,
            "cmd": cmd,
            "returncode": result.returncode,
            "traceback": "Subprocess failed",
            "stderr": stderr_msg,
        })
        print(f"[FAILURE CACHED] {task_id} - Reason: {stderr_msg.strip()[:100]}...")



def cleanup(local_task_dir: str):
    shutil.rmtree(local_task_dir, ignore_errors=True)


def main():
    cache = diskcache.Cache(CACHE_DIR)
    os.makedirs(LOCAL_TMP_ROOT, exist_ok=True)

    wb = load_workbook(XLSX_IN)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    col_idx = {str(name).strip(): i for i, name in enumerate(headers) if name is not None}

    required_cols = [TASK_ID_COLUMN, TASK_DSCP_COLUMN, DEVICE_COLUMN]
    for c in required_cols:
        if c not in col_idx:
            raise KeyError(f"Excel missing required column: {c}")

    task_id_col = col_idx[TASK_ID_COLUMN]
    task_dscp_col = col_idx[TASK_DSCP_COLUMN]
    device_col = col_idx[DEVICE_COLUMN]

    task_rows = []
    seen_task_ids = set()
    for row in ws.iter_rows(min_row=2):
        cell_val = row[task_id_col].value
        if cell_val is None:
            continue
        task_id = str(cell_val).strip()
        if not task_id:
            continue
        if task_id in seen_task_ids:
            continue

        # Filter for specific target tasks
        if TARGET_TASKS and task_id not in TARGET_TASKS:
            continue

        device_name = row[device_col].value
        device_name = str(device_name).strip() if device_name is not None else ""
        if device_name not in ALLOWED_DEVICES:
            continue

        seen_task_ids.add(task_id)
        task_dscp = row[task_dscp_col].value
        task_rows.append((task_id, task_dscp, device_name))

    ready_q = Queue()
    slots = Queue(maxsize=PARALLEL_CVT_TASKS + MAX_AHEAD)

    def cvt_and_cleanup(task_id: str, local_dir: str, task_dscp: str, device_name: str, cvt_threads: int):
        try:
            cvt_task(task_id, local_dir, task_dscp, device_name, workers=cvt_threads, cache=cache)
            cache.set("cached_task_id", task_id)
            cleanup(local_dir)
            print(f"[DONE] CVT {task_id}")
        except Exception:
            print(traceback.format_exc())
            os._exit(1)
        finally:
            slots.get()
            slots.task_done()
        return task_id

    def cvt_submitter(cache_obj):
        while True:
            item = ready_q.get()
            try:
                if item is SENTINEL:
                    return
                task_id, local_dir, task_dscp, device_name = item
                print("cvt get", task_id, device_name)
                cvt_and_cleanup(task_id, local_dir, task_dscp, device_name, CVT_WORKERS_PER_TASK)
            finally:
                ready_q.task_done()

    def submit_download(dl_pool: ThreadPoolExecutor, task_id: str, task_dscp: str, device_name: str):
        def _job():
            local_dir = os.path.join(LOCAL_TMP_ROOT, task_id)
            print(f"\n=== DOWNLOAD TASK {task_id} ({device_name}) ===")
            download_task(task_id, local_dir, workers=DOWNLOAD_WORKERS_PER_TASK)
            print(f"\n=== DOWNLOAD TASK FINISH {task_id} ({device_name}) ===")
            ready_q.put((task_id, local_dir, task_dscp, device_name))
        return dl_pool.submit(_job)

    cached_task_id = cache.get("cached_task_id")
    skip_mode = cached_task_id is not None
    seen_cached = False

    with ThreadPoolExecutor(max_workers=PARALLEL_DOWNLOAD_TASKS) as dl_pool:
        t = threading.Thread(target=cvt_submitter, args=(cache,), daemon=True)
        t.start()

        download_futures = []

        for task_id, task_dscp, device_name in tqdm(task_rows, total=len(task_rows), miniters=1):
            if skip_mode:
                if not seen_cached:
                    if task_id == cached_task_id:
                        seen_cached = True
                    continue
                else:
                    if task_id == cached_task_id:
                        continue
                    skip_mode = False
                    cached_task_id = None

            print("ready_q size:", ready_q.qsize())
            print("slots size:", slots.qsize())

            slots.put(1)
            download_futures.append(submit_download(dl_pool, task_id, task_dscp, device_name))

        for fut in as_completed(download_futures):
            fut.result()

        ready_q.join()
        ready_q.put(SENTINEL)
        t.join()

    wb.close()


if __name__ == "__main__":
    main()
