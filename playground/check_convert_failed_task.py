# -*- coding: utf-8 -*-
import os
from openpyxl import load_workbook

# ===================== TOP CONFIG =====================

HOME = "/root/jyw_ws/lerobot_jpeg/playground"
XLSX_IN = f"{HOME}/数据需求总表同步_转换列表.xlsx"
OUTPUT_ROOT = "/mnt/sdc/output_0331"

TASK_ID_COLUMN = "任务ID"
DEVICE_COLUMN = "采集设备"

DEVICE_TO_ROBOT = {
    "松灵Aloha": "cobotmagic",
    "Franka FR3": "fr3",
    "傅利叶GR-2": "gr2",
    "星尘智能S1": "astribot_s1",
}
ALLOWED_DEVICES = set(DEVICE_TO_ROBOT.keys())

# ======================================================

def is_task_done(task_id: str) -> bool:
    """
    判断任务是否已在输出目录中存在。
    这里只检查 OUTPUT_ROOT/task_id 是否存在。
    如需更严格，可改成检查是否含有 meta / data 等文件。
    """
    out_dir = os.path.join(OUTPUT_ROOT, task_id)
    return os.path.isdir(out_dir)

def main():
    wb = load_workbook(XLSX_IN)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    col_idx = {str(name).strip(): i for i, name in enumerate(headers) if name is not None}
 

    required_cols = [TASK_ID_COLUMN, DEVICE_COLUMN]
    for c in required_cols:
        if c not in col_idx:
            raise KeyError(f"Excel missing required column: {c}")

    task_id_col = col_idx[TASK_ID_COLUMN]
    device_col = col_idx[DEVICE_COLUMN]

    seen_task_ids = set()
    all_valid_task_ids = []

    for row in ws.iter_rows(min_row=2):
        cell_val = row[task_id_col].value
        if cell_val is None:
            continue

        task_id = str(cell_val).strip()
        if not task_id or task_id in seen_task_ids:
            continue

        device_name = row[device_col].value
        device_name = str(device_name).strip() if device_name is not None else ""
        if device_name not in ALLOWED_DEVICES:
            continue

        seen_task_ids.add(task_id)
        all_valid_task_ids.append((task_id,device_name))

    wb.close()

    missing_tasks = [task_id for task_id in all_valid_task_ids if not is_task_done(task_id[0])]

    print("=" * 80)
    print(f"total valid tasks in xlsx : {len(all_valid_task_ids)}")
    print(f"missing tasks in output   : {len(missing_tasks)}")
    print("=" * 80)

    for task_id in missing_tasks:
        print(task_id)

    with open("missing_tasks.txt", "w", encoding="utf-8") as f:
        for task_id in missing_tasks:
            f.write(task_id + "\n")

    print("=" * 80)
    print("saved to missing_tasks.txt")

if __name__ == "__main__":
    main()