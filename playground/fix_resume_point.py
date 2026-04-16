# -*- coding: utf-8 -*-
import os
import diskcache
from openpyxl import load_workbook

# ===================== CONFIG (Must match convert.new2.py) =====================
HOME = "/root/jyw_ws/lerobot_jpeg/playground"
XLSX_IN = f"{HOME}/数据需求总表同步_转换列表.xlsx"
CACHE_DIR = "/root/_robot2_ckpt"

TASK_ID_COLUMN = "任务ID"
DEVICE_COLUMN = "采集设备"
DEVICE_TO_ROBOT = {"傅利叶GR-2": "gr2", "星尘智能S1": "astribot_s1"}
ALLOWED_DEVICES = set(DEVICE_TO_ROBOT.keys())
# ==============================================================================

def get_filtered_tasks():
    wb = load_workbook(XLSX_IN)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    col_idx = {str(name).strip(): i for i, name in enumerate(headers) if name is not None}
    task_id_col, device_col = col_idx[TASK_ID_COLUMN], col_idx[DEVICE_COLUMN]
    
    task_rows = []
    seen = set()
    for row in ws.iter_rows(min_row=2):
        tid = str(row[task_id_col].value).strip() if row[task_id_col].value else None
        if not tid or tid in seen: continue
        dev = str(row[device_col].value).strip() if row[device_col].value else ""
        if dev not in ALLOWED_DEVICES: continue
        seen.add(tid)
        task_rows.append(tid)
    wb.close()
    return task_rows

def main():
    cache = diskcache.Cache(CACHE_DIR)
    tasks = get_filtered_tasks()
    
    # The ID from your diagnostic: Last actual success was Index 157
    CORRECT_LAST_SUCCESS = "03a90d45c38b4a23b61e4ae7759237c4"
    
    if CORRECT_LAST_SUCCESS not in tasks:
        print(f"❌ Error: {CORRECT_LAST_SUCCESS} not found in Excel list!")
        return

    # 1. Reset the resume point
    old_id = cache.get("cached_task_id")
    cache.set("cached_task_id", CORRECT_LAST_SUCCESS)
    print(f"✅ Updated cached_id: {old_id} -> {CORRECT_LAST_SUCCESS}")

    # 2. Clear failure cache for all subsequent tasks
    start_idx = tasks.index(CORRECT_LAST_SUCCESS)
    tasks_to_clear = tasks[start_idx + 1:]
    
    cleared_count = 0
    for tid in tasks_to_clear:
        failure_key = f"convert_failure_{tid}"
        if failure_key in cache:
            cache.delete(failure_key)
            cleared_count += 1
            
    print(f"✅ Cleared {cleared_count} old failure records from the cache.")
    print(f"\nResume state is now clean. Process will start from Index {start_idx + 1}.")
    print("You can now run: python convert.new2.py")

if __name__ == "__main__":
    main()