# -*- coding: utf-8 -*-
import os
import shutil
import subprocess
import threading
import traceback
from queue import Queue
from concurrent.futures import ThreadPoolExecutor, as_completed

import diskcache
from openpyxl import load_workbook
from tqdm import tqdm

# ===================== TOP CONFIG =====================

# 并行 task 数
PARALLEL_DOWNLOAD_TASKS = 1
PARALLEL_CVT_TASKS = 1

# 每个 task 内部资源
DOWNLOAD_WORKERS_PER_TASK = 4
CVT_THREADS_PER_TASK = 8

# 背压：最多下载领先多少个 task
MAX_AHEAD = 5
SENTINEL = object()

# 路径相关
HOME = "/root/lerobot_jpeg/playground"


RCLONE_BIN = f"{HOME}/rclone-v1.68.2-linux-amd64/rclone"
XLSX_IN = f"{HOME}/gr2data0331.xlsx"
CVT_SCRIPT = f"{HOME}/gr2_align2lerobot_v30.jpeg.py"

LOCAL_TMP_ROOT = "/mnt/fastdisk/tmp_download"
OUTPUT_ROOT = "/mnt/sdc/output_0331"
CACHE_DIR = "/root/_gr2_ckpt2"

# 远端相关
REMOTE_ROOT = "huawei-cloud:openloong-apps-prod-private/data-collector-svc/align"

# Excel 列名
TASK_ID_COLUMN = "任务ID"
TASK_DSCP_COLUMN = "任务步骤描述-人话英文版"

# ======================================================


def run(cmd, check=True):
    p = subprocess.Popen(cmd, start_new_session=False)
    ret = p.wait()
    if check and ret != 0:
        raise subprocess.CalledProcessError(ret, cmd)


def download_task(task_id: str, local_task_dir: str, workers: int):
    remote_task_root = f"{REMOTE_ROOT}/{task_id}"
    os.makedirs(local_task_dir, exist_ok=True)

    cmd = [
        RCLONE_BIN, "copy",
        remote_task_root,
        local_task_dir,
        "--transfers", str(workers),
        "--progress"
    ]
    run(cmd, check=True)


def cvt_task(task_id: str, local_task_dir: str, task_dscp: str, workers: int):
    output_dir = f"{OUTPUT_ROOT}/{task_id}"
    cmd = [
        "python", CVT_SCRIPT,
        "--input", local_task_dir,
        "--output", output_dir,
        "--repo_id", task_id,
        "--task", str(task_dscp) if task_dscp is not None else "",
        "--workers", str(workers),
    ]
    print("task_dscp", task_dscp)
    run(cmd, check=True)


def cleanup(local_task_dir: str):
    shutil.rmtree(local_task_dir, ignore_errors=True)


def main():
    cache = diskcache.Cache(CACHE_DIR)
    os.makedirs(LOCAL_TMP_ROOT, exist_ok=True)

    wb = load_workbook(XLSX_IN)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    col_idx = {str(name).strip(): i for i, name in enumerate(headers) if name is not None}

    required_cols = [TASK_ID_COLUMN, TASK_DSCP_COLUMN]
    for c in required_cols:
        if c not in col_idx:
            raise KeyError(f"Excel missing required column: {c}")

    task_id_col = col_idx[TASK_ID_COLUMN]
    task_dscp_col = col_idx[TASK_DSCP_COLUMN]

    task_rows = []
    seen_task_ids = set()
    for row in ws.iter_rows(min_row=2):
        cell_val = row[task_id_col].value
        if cell_val is None:
            continue
        task_id = str(cell_val).strip()
        if not task_id:
            continue
        if task_id in seen_task_ids:
            continue
        seen_task_ids.add(task_id)

        task_dscp = row[task_dscp_col].value
        task_rows.append((task_id, task_dscp))

    ready_q = Queue()
    slots = Queue(maxsize=PARALLEL_CVT_TASKS + MAX_AHEAD)

    def cvt_and_cleanup(task_id: str, local_dir: str, task_dscp: str, cvt_threads: int):
        try:
            cvt_task(task_id, local_dir, task_dscp, workers=cvt_threads)
            cache.set("cached_task_id", task_id)
            cleanup(local_dir)
            print(f"[DONE] CVT {task_id}")
        except Exception:
            print(traceback.format_exc())
            os._exit(1)
        finally:
            slots.get()
            slots.task_done()
        return task_id

    def cvt_submitter(cache_obj):
        while True:
            item = ready_q.get()
            try:
                if item is SENTINEL:
                    return
                task_id, local_dir, task_dscp = item
                print("cvt get", task_id)
                cvt_and_cleanup(task_id, local_dir, task_dscp, CVT_THREADS_PER_TASK)
            finally:
                ready_q.task_done()

    def submit_download(dl_pool: ThreadPoolExecutor, task_id: str, task_dscp: str):
        def _job():
            local_dir = os.path.join(LOCAL_TMP_ROOT, task_id)
            print(f"\n=== DOWNLOAD TASK {task_id} ===")
            download_task(task_id, local_dir, workers=DOWNLOAD_WORKERS_PER_TASK)
            print(f"\n=== DOWNLOAD TASK FINISH {task_id} ===")
            ready_q.put((task_id, local_dir, task_dscp))
        return dl_pool.submit(_job)

    cached_task_id = cache.get("cached_task_id")
    skip_mode = cached_task_id is not None
    seen_cached = False

    with ThreadPoolExecutor(max_workers=PARALLEL_DOWNLOAD_TASKS) as dl_pool:
        t = threading.Thread(target=cvt_submitter, args=(cache,), daemon=True)
        t.start()

        download_futures = []

        for task_id, task_dscp in tqdm(task_rows, total=len(task_rows), miniters=1):
            if skip_mode:
                if not seen_cached:
                    if task_id == cached_task_id:
                        seen_cached = True
                    continue
                else:
                    if task_id == cached_task_id:
                        continue
                    skip_mode = False
                    cached_task_id = None

            print("ready_q size:", ready_q.qsize())
            print("slots size:", slots.qsize())

            slots.put(1)
            download_futures.append(submit_download(dl_pool, task_id, task_dscp))

        for fut in as_completed(download_futures):
            fut.result()

        ready_q.join()
        ready_q.put(SENTINEL)
        t.join()

    wb.close()


if __name__ == "__main__":
    main()