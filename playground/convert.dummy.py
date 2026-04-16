# -*- coding: utf-8 -*-
"""
convert.dummy.py - Task Conversion Status Tracker

Purpose:
    This script tracks and identifies the last successfully converted task 
    in a data conversion pipeline. It reads task information from an Excel 
    file, filters by allowed device types, and checks output directories 
    to determine conversion status.

Workflow:
    1. Load task rows from Excel (filtered by device type)
    2. Iterate through tasks in Excel order
    3. Check if each task's output directory exists and is non-empty
    4. Cache the last converted task ID for resume capability

Usage:
    - Run to find the last converted task ID
    - Result is cached in diskcache for workflow resumption
    - Useful for resuming interrupted batch conversion processes

Configuration:
    - XLSX_IN: Path to the task list Excel file
    - OUTPUT_ROOT: Root directory where converted task outputs are stored
    - DEVICE_TO_ROBOT: Mapping of device names to robot identifiers
    - CACHE_DIR: Location for diskcache storage

Author: [Your Name]
Date: 2024-03-31
"""

import os
from openpyxl import load_workbook

# ===================== SAME CONFIG =====================
HOME = "/root/jyw_ws/lerobot_jpeg/playground"
XLSX_IN = f"{HOME}/数据需求总表同步_转换列表.xlsx"
OUTPUT_ROOT = "/mnt/sdc/output_0331"

TASK_ID_COLUMN = "任务ID"
TASK_DSCP_COLUMN = "任务步骤描述-人话英文版"
DEVICE_COLUMN = "采集设备"

DEVICE_TO_ROBOT = {
    "松灵Aloha": "cobotmagic",
    "Franka FR3": "fr3",
}

ALLOWED_DEVICES = set(DEVICE_TO_ROBOT.keys())
# ======================================================

def build_task_rows():
    wb = load_workbook(XLSX_IN)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    col_idx = {str(name).strip(): i for i, name in enumerate(headers) if name is not None}

    required_cols = [TASK_ID_COLUMN, TASK_DSCP_COLUMN, DEVICE_COLUMN]
    for c in required_cols:
        if c not in col_idx:
            raise KeyError(f"Excel missing required column: {c}")

    task_id_col = col_idx[TASK_ID_COLUMN]
    task_dscp_col = col_idx[TASK_DSCP_COLUMN]
    device_col = col_idx[DEVICE_COLUMN]

    task_rows = []
    seen_task_ids = set()

    for row in ws.iter_rows(min_row=2):
        cell_val = row[task_id_col].value
        if cell_val is None:
            continue

        task_id = str(cell_val).strip()
        if not task_id:
            continue

        # 和原逻辑一致：任务ID去重，保留第一次出现
        if task_id in seen_task_ids:
            continue

        device_name = row[device_col].value
        device_name = str(device_name).strip() if device_name is not None else ""

        # 和原逻辑一致：只保留允许设备
        if device_name not in ALLOWED_DEVICES:
            continue

        seen_task_ids.add(task_id)
        task_dscp = row[task_dscp_col].value
        task_rows.append((task_id, task_dscp, device_name))

    wb.close()
    return task_rows


def is_task_converted(task_id: str) -> bool:
    """
    尽量贴近原逻辑：
    原脚本成功后 output_dir = f"{OUTPUT_ROOT}/{task_id}"
    所以这里判断该目录是否存在，且里面有实际输出内容。
    """
    output_dir = os.path.join(OUTPUT_ROOT, task_id)
    if not os.path.isdir(output_dir):
        return False

    # 比单纯 exists 更稳一点：目录非空
    try:
        return len(os.listdir(output_dir)) > 0
    except Exception:
        return False


def find_last_converted_task_id_by_same_logic():
    task_rows = build_task_rows()

    last_converted = None
    last_index = None

    for i, (task_id, task_dscp, device_name) in enumerate(task_rows):
        if is_task_converted(task_id):
            last_converted = task_id
            last_index = i

    if last_converted is None:
        print("No converted task found under OUTPUT_ROOT using same Excel/filter order.")
        return None

    print("=" * 100)
    print("last converted task id:", last_converted)
    print("index in filtered task_rows:", last_index)
    print("=" * 100)

    # 顺便打印前后几个，方便核对恢复点
    start = max(0, last_index - 3)
    end = min(len(task_rows), last_index + 4)

    print("\nNearby tasks in same logic order:")
    for i in range(start, end):
        task_id, task_dscp, device_name = task_rows[i]
        mark = "<<< LAST_CONVERTED" if i == last_index else ""
        converted = is_task_converted(task_id)
        print(f"[{i}] task_id={task_id}  converted={converted}  device={device_name} {mark}")

    return last_converted


if __name__ == "__main__":
    # find_last_converted_task_id_by_same_logic()

    # -*- coding: utf-8 -*-
    import diskcache

    CACHE_DIR = "/root/_robot_ckpt"

    last_task_id = find_last_converted_task_id_by_same_logic()
    if last_task_id is not None:
        cache = diskcache.Cache(CACHE_DIR)
        cache.set("cached_task_id", last_task_id)
        print(f"\nRestored cache: cached_task_id = {last_task_id}")