# -*- coding: utf-8 -*-
from collections import defaultdict
from openpyxl import load_workbook

HOME = "/root/jyw_ws/lerobot_jpeg/playground"
XLSX_IN = f"{HOME}/数据需求总表同步_转换列表.xlsx"

TASK_ID_COLUMN = "任务ID"
TASK_DSCP_COLUMN = "任务步骤描述-人话英文版"
DEVICE_COLUMN = "采集设备"
HOUR_COLUMN = "采集时长（小时）"

TARGET_DEVICES = {
    "星尘智能S1",
    "Franka FR3",
    "傅利叶GR-2",
    "松灵Aloha",
}

TASKS_PER_DEVICE = 5

def to_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    try:
        return float(s)
    except Exception:
        return None


def main():
    wb = load_workbook(XLSX_IN, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    col_idx = {str(name).strip(): i for i, name in enumerate(headers) if name is not None}

    required_cols = [TASK_ID_COLUMN, DEVICE_COLUMN, HOUR_COLUMN]
    for c in required_cols:
        if c not in col_idx:
            raise KeyError(f"Excel missing required column: {c}")

    task_id_col = col_idx[TASK_ID_COLUMN]
    device_col = col_idx[DEVICE_COLUMN]
    hour_col = col_idx[HOUR_COLUMN]
    task_dscp_col = col_idx.get(TASK_DSCP_COLUMN, None)

    # key = 采集设备原值
    robot_tasks = defaultdict(list)
    seen_task_ids = set()

    for row in ws.iter_rows(min_row=2):
        task_id_val = row[task_id_col].value
        if task_id_val is None:
            continue

        task_id = str(task_id_val).strip()
        if not task_id or task_id in seen_task_ids:
            continue

        device_name = row[device_col].value
        device_name = str(device_name).strip() if device_name is not None else ""
        if not device_name:
            continue

        if device_name not in TARGET_DEVICES:
            continue


        hour_val = to_float(row[hour_col].value)
        if hour_val is None or hour_val <= 0:
            continue

        task_dscp = row[task_dscp_col].value if task_dscp_col is not None else ""

        seen_task_ids.add(task_id)
        robot_tasks[device_name].append({
            "task_id": task_id,
            "hour": hour_val,
            "device_name": device_name,
            "task_dscp": task_dscp,
        })

    print("=" * 120)
    print("TEST SET TASK PICKER - group by 采集设备, sort by 采集时长（小时）")
    print("=" * 120)

    for device_name, items in robot_tasks.items():

        selected = sorted(items, key=lambda x: (x["hour"], x["task_id"]))[:TASKS_PER_DEVICE]

        print("\n" + "=" * 120)
        print(f"DEVICE: {device_name}")
        print("=" * 120)

        print(f"{'#':<5} | {'Task ID':<36} | {'Hour':<10} | Task Description")
        print("-" * 120)
        for i, x in enumerate(selected, 1):
            print(f"{i:<5} | {x['task_id']:<36} | {x['hour']:<10.2f} | {x['task_dscp']}")

    print("\n" + "=" * 120)
    print("SUMMARY")
    print("=" * 120)
    for device_name, items in robot_tasks.items():

        print(f"{device_name:<20} | total={len(items):<5} | selected={min(len(items), TASKS_PER_DEVICE)}")

    wb.close()


if __name__ == "__main__":
    main()